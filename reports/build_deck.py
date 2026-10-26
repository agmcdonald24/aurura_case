#!/usr/bin/env python
"""
Build the take-home deck as a PowerPoint file from reports/data/ and reports/figures/.

Usage, from the repo root:

    python reports/pull_deck_data.py      # on the Mac: runs the queries, writes data/
    python reports/build_charts.py        # anywhere: writes figures/
    python reports/build_deck.py          # anywhere: writes Aurora_Construction_Take_Home.pptx

Needs python-pptx. The deck is an output; this script is the source. A number that
changes means re-pulling and re-running, never editing the slides by hand. Every slide's
footer names the query its numbers came from.

The deck walks through the work in the order it happened: environment, a first look at
the data, the problems found, the staging rules that came out of them, then the four
questions, then how the answers would be delivered to two audiences and how AI was used.
"""

import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu
from lxml import etree

HERE = Path(__file__).resolve().parent
DATA = HERE / "data"
FIGURES = HERE / "figures"
MEDIA = HERE / "media"
OUT = HERE / "Aurora_Construction_Take_Home.pptx"

# --- brand -------------------------------------------------------------------
BLUE = RGBColor(0x13, 0x68, 0xE5)      # Aurora blue, sampled from the wordmark
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x7B, 0x8C)
RULE = RGBColor(0xDC, 0xE3, 0xEA)
LIGHT = RGBColor(0xF1, 0xF4, 0xF8)
HI = RGBColor(0xE6, 0xF0, 0xFD)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
TITLE_Y, TITLE_H = 0.42, 0.9
BODY_Y = 1.45
FOOTER_Y = 6.82

NO_GRID_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # "No Style, No Grid"
REPO = "https://github.com/agmcdonald24/aurura_case/blob/main/"


def link(path: str) -> str:
    return REPO + path


# --- data --------------------------------------------------------------------
def read(name: str) -> list[dict]:
    with open(DATA / f"{name}.csv", newline="") as f:
        return list(csv.DictReader(f))


def package() -> dict:
    p = DATA / "package.json"
    return json.loads(p.read_text()) if p.exists() else {}


def n(x) -> float:
    return float(x) if x not in ("", None) else 0.0


def millions(x) -> str:
    v = n(x) / 1e6
    return f"{v:.1f}M" if v < 100 else f"{v:.0f}M"


def thousands(x) -> str:
    return f"{n(x):,.0f}"


def compact(x) -> str:
    v = n(x)
    if v >= 1e6:
        return millions(v)
    if v >= 1e4:
        return f"{v/1e3:.0f}k"
    return f"{v:,.0f}"


# --- primitives --------------------------------------------------------------
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        self.pulled = package().get("pulled_at", "")[:10]
        self.count = 0

    # -- text ----------------------------------------------------------------
    def text(self, slide, x, y, w, h, paragraphs, size=14, color=INK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.1, space_after=4):
        """paragraphs: list of str, or list of lists of (text, {bold, color, size, italic})."""
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        first = True
        for para in paragraphs:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align
            p.line_spacing = line_spacing
            p.space_after = Pt(space_after)
            runs = [(para, {})] if isinstance(para, str) else para
            for txt, opts in runs:
                r = p.add_run()
                r.text = txt
                f = r.font
                f.name = FONT
                f.size = Pt(opts.get("size", size))
                f.bold = opts.get("bold", bold)
                f.italic = opts.get("italic", False)
                f.color.rgb = opts.get("color", color)
                if opts.get("link"):
                    r.hyperlink.address = opts["link"]
                    f.color.rgb = BLUE
                    f.underline = True
        return box

    def prose(self, slide, x, y, w, h, paragraphs, size=14, color=INK):
        """Plain paragraphs, the way it would be said out loud."""
        return self.text(slide, x, y, w, h, paragraphs, size=size, color=color,
                         line_spacing=1.18, space_after=10)

    def bullets(self, slide, x, y, w, h, items, size=14, color=INK, space_after=6):
        """Plain bullet list. Each item is a str or a list of (text, opts) runs."""
        box = self.text(slide, x, y, w, h, items, size=size, color=color, space_after=space_after)
        for p in box.text_frame.paragraphs:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.22))))
            pPr.set("indent", str(-Emu(Inches(0.22))))
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")
        return box

    # -- shapes --------------------------------------------------------------
    def rect(self, slide, x, y, w, h, fill=LIGHT, line=None, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = fill
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
            s.line.width = Pt(0.75)
        s.shadow.inherit = False
        s.text_frame.text = ""
        return s

    def hline(self, slide, x, y, w, color=RULE, weight=0.75):
        ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
        ln.line.color.rgb = color
        ln.line.width = Pt(weight)
        return ln

    def card(self, slide, x, y, w, h, heading, number, lines, size=12.5):
        self.rect(slide, x, y, w, h, fill=LIGHT)
        self.text(slide, x + 0.25, y + 0.2, w - 0.5, 0.4, [heading], size=14, bold=True)
        self.text(slide, x + 0.25, y + 0.6, w - 0.5, 0.7, [number], size=30, bold=True, color=BLUE)
        self.text(slide, x + 0.25, y + 1.35, w - 0.5, h - 1.5, lines, size=size, space_after=5)

    def stat(self, slide, x, y, w, number, label, size=34):
        self.text(slide, x, y, w, 0.65, [number], size=size, bold=True, color=BLUE)
        self.text(slide, x, y + 0.62, w, 0.5, [label], size=12, color=MUTED)

    # -- tables --------------------------------------------------------------
    def table(self, slide, x, y, widths, rows, size=11.5, row_h=0.32, highlight=None,
              align_right_from=1, header_size=None, header=True):
        """rows[0] is the header unless header=False. highlight: set of row indexes to tint."""
        highlight = highlight or set()
        n_rows, n_cols = len(rows), len(rows[0])
        shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y),
                                       Inches(sum(widths)), Inches(row_h * n_rows))
        tbl = shape.table
        tblPr = tbl._tbl.tblPr
        style = tblPr.find(qn("a:tableStyleId"))
        if style is None:
            style = etree.SubElement(tblPr, qn("a:tableStyleId"))
        style.text = NO_GRID_STYLE
        tbl.first_row = False
        tbl.horz_banding = False
        for j, w in enumerate(widths):
            tbl.columns[j].width = Inches(w)
        for i in range(n_rows):
            tbl.rows[i].height = Inches(row_h)
            for j in range(n_cols):
                cell = tbl.cell(i, j)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT if j >= align_right_from else PP_ALIGN.LEFT
                r = p.add_run()
                r.text = str(rows[i][j])
                r.font.name = FONT
                is_header = header and i == 0
                r.font.size = Pt(header_size or size) if is_header else Pt(size)
                r.font.bold = is_header or (i in highlight) or (not header and j == 0)
                r.font.color.rgb = MUTED if is_header else INK
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT if is_header else (HI if i in highlight else WHITE)
                if is_header:
                    self._border(cell, "a:lnB", RULE, 1.0)
        return shape

    @staticmethod
    def _border(cell, side, color, width_pt):
        tcPr = cell._tc.get_or_add_tcPr()
        ln = etree.SubElement(tcPr, qn(side))
        ln.set("w", str(int(width_pt * 12700)))
        solid = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(solid, qn("a:srgbClr"))
        clr.set("val", "%02X%02X%02X" % (color[0], color[1], color[2]))

    def caption(self, slide, x, y, w, text):
        self.text(slide, x, y, w, 0.35, [text], size=10, color=MUTED)

    # -- slide frames --------------------------------------------------------
    def content_slide(self, title, takeaway=None, source=None):
        slide = self.prs.slides.add_slide(self.blank)
        self.count += 1
        self.text(slide, MARGIN, TITLE_Y, SLIDE_W - 2 * MARGIN - 0.9, TITLE_H, [title],
                  size=26, bold=True, anchor=MSO_ANCHOR.TOP, line_spacing=1.0)
        mark = MEDIA / "aurora_mark.png"
        if mark.exists():
            slide.shapes.add_picture(str(mark), Inches(SLIDE_W - MARGIN - 0.42), Inches(0.45),
                                     height=Inches(0.46))
        self.hline(slide, MARGIN, FOOTER_Y - 0.08, SLIDE_W - 2 * MARGIN)
        if takeaway:
            self.text(slide, MARGIN, FOOTER_Y, 6.9, 0.55, [takeaway], size=12, color=INK)
        runs = []
        if source:
            runs.append(("Source: ", {}))
            for i, (label, path) in enumerate(source):
                if i:
                    runs.append((", ", {}))
                runs.append((label, {"link": link(path)}))
            if self.pulled:
                runs.append((f"   ·   data pulled {self.pulled}", {}))
        runs.append((f"   ·   {self.count + 1}", {}))
        self.text(slide, SLIDE_W - MARGIN - 5.2, FOOTER_Y, 5.2, 0.55, [runs],
                  size=9, color=MUTED, align=PP_ALIGN.RIGHT)
        return slide

    def cover(self, title, subtitle, byline):
        slide = self.prs.slides.add_slide(self.blank)
        panel_w = 5.3
        photo = MEDIA / "aurora_truck.jpg"
        if photo.exists():
            pic = slide.shapes.add_picture(str(photo), Inches(panel_w), 0,
                                           Inches(SLIDE_W - panel_w), Inches(SLIDE_H))
            src_aspect = 2576 / 1449
            box_aspect = (SLIDE_W - panel_w) / SLIDE_H
            visible = box_aspect / src_aspect          # fraction of source width shown
            pic.crop_left = pic.crop_right = (1 - visible) / 2
        wordmark = MEDIA / "aurora_wordmark.png"
        if wordmark.exists():
            slide.shapes.add_picture(str(wordmark), Inches(MARGIN), Inches(0.75), width=Inches(2.3))
        self.text(slide, MARGIN, 2.4, panel_w - MARGIN - 0.3, 1.0, [title], size=36, bold=True,
                  line_spacing=1.0)
        self.text(slide, MARGIN, 3.45, panel_w - MARGIN - 0.3, 0.6, [subtitle[0]], size=17, color=BLUE,
                  bold=True, line_spacing=1.1)
        self.text(slide, MARGIN, 4.2, panel_w - MARGIN - 0.3, 1.8, [subtitle[1]], size=14, color=INK,
                  line_spacing=1.2)
        self.text(slide, MARGIN, 6.55, panel_w - MARGIN - 0.3, 0.6, [byline], size=11, color=MUTED)
        return slide

    def figure(self, slide, name, x, y, w=None, h=None):
        path = FIGURES / name
        kwargs = {}
        if w:
            kwargs["width"] = Inches(w)
        if h:
            kwargs["height"] = Inches(h)
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)

    def save(self):
        self.prs.save(OUT)
        print(f"wrote {OUT.relative_to(HERE.parent)} ({self.count + 1} slides)")


# --- the slides --------------------------------------------------------------
def build():
    d = Deck()
    counts = {r["measure"]: n(r["value"]) for r in read("12_row_counts")}
    flags = {r["event_type"]: r for r in read("10_staging_flags")}
    con = flags["construction"]
    top = read("01_q1_top_cities")[:10]
    first, second = top[0], top[1]
    cmp_rows = read("07_q3_city_comparison")
    routes_all = [r for r in read("08_q4_route_exposure") if r["year"] == "all"]
    routes_year = [r for r in read("08_q4_route_exposure") if r["year"] != "all"]
    decomp = read("05_q2_decomposition_by_year")
    by_year = read("03_q2_ten_cities_by_year")
    dup = counts["duplicates removed"]
    desc = read("09_q4_berkeley_2021_descriptions")
    examples = read("13_seasonal_pattern_examples")
    bridges = read("route_bridges")

    def route_name(r):
        return r["path_name"].replace("SF -> ", "").replace("Berkley", "Berkeley")

    L = {  # repo paths the slides link to
        "loader": "scripts/load_raw.py",
        "staging": "aurora_dbt/models/staging/stg_construction.sql",
        "sources": "aurora_dbt/models/staging/_sources.yml",
        "seeds": "aurora_dbt/seeds/_seeds.yml",
        "int": "aurora_dbt/models/intermediate/int_city_month_hours.sql",
        "routes_stg": "aurora_dbt/models/staging/stg_routes.sql",
        "routes_int": "aurora_dbt/models/intermediate/int_route_events.sql",
        "routes_fct": "aurora_dbt/models/marts/fct_route_exposure.sql",
        "cmp_fct": "aurora_dbt/models/marts/fct_city_comparison.sql",
        "dict": "docs/data_dictionary.md",
        "log": "process.md",
        "bridges": "reports/check_route_bridges.py",
        "deck": "reports/build_deck.py",
    }
    Q = {i: f"reports/queries/{name}" for i, name in enumerate([
        "01_q1_top_cities.sql", "02_q1_ten_cities.sql", "03_q2_ten_cities_by_year.sql",
        "04_q2_national_monthly.sql", "05_q2_decomposition_by_year.sql", "06_q2_new_york_by_year.sql",
        "07_q3_city_comparison.sql", "08_q4_route_exposure.sql", "09_q4_berkeley_2021_descriptions.sql",
        "10_staging_flags.sql", "11_duration_distribution.sql", "12_row_counts.sql",
        "13_seasonal_pattern_examples.sql"], start=1)}

    # 1 — cover
    d.cover("Construction, counted",
            ("Data Analyst take-home",
             "Six questions about road construction, answered from a public dataset of "
             "six million events, with the work shown in the order it happened."),
            "Andrew McDonald   ·   October 2026   ·   US Road Construction and Closures, 2016–2021")

    # 2 — the questions
    s = d.content_slide("The questions", source=[("repo", "README.md"), ("process log", L["log"])])
    d.prose(s, MARGIN, BODY_Y, 7.6, 4.8, [
        "1.  Over the past five years, which city has had the most hours of construction?",
        "2.  How has this changed over time?",
        "3.  Accounting for varying city sizes, how do the ten largest US cities compare?",
        "4.  Of five potential self-driving routes, which poses the greatest risk of autonomy "
        "challenges because of construction?",
        "5.  How would I present that answer differently to engineers and to executives?",
        "6.  What AI-based tools did I use, how, and why?",
    ], size=14)
    d.prose(s, 8.7, BODY_Y, 4.0, 4.8, [
        "This deck goes through what I did in the order I did it: setting up, looking at the "
        "data, fixing what I found, then the questions.",
        [("All of the code is in the repo, and each slide links to the model or query behind it. "
          "The log I kept along the way is ", {}), ("process.md", {"link": link(L["log"])}), (".", {})],
    ], size=13)

    # 3 — tools
    s = d.content_slide("The tools I used", source=[("loader", L["loader"]), ("query 12", Q[12])])
    d.prose(s, MARGIN, BODY_Y, 7.4, 3.6, [
        "Opening up the dataset and seeing over six million rows to pull in, I decided to spin up "
        "a Postgres instance on my local machine. I created a Python virtual environment connected "
        "to a GitHub repository and installed dbt as a way to transform and work with the data. "
        "From there I loaded the file into Postgres so I could play with it and manipulate it in "
        "a way that is easy and familiar to me.",
        [("The load itself was a ", {}), ("short Python script", {"link": link(L["loader"])}),
         (" that creates the table with proper column types and streams the file in. If a value "
          "had not fit its type the load would have failed, and it didn't, so that was my first "
          "check. The five routes are a tiny file, so they went into dbt as a ", {}),
         ("seed", {"link": link(L["seeds"])}), (".", {})],
        [("Everything downstream is a dbt model: a ", {}), ("staging layer", {"link": link(L["staging"])}),
         (" that cleans the data, an ", {}), ("intermediate layer", {"link": link(L["int"])}),
         (" that spreads each event's hours across the months it was active, and the final "
          "tables that answer the questions.", {})],
    ], size=13.5)
    d.stat(s, 8.7, BODY_Y + 0.1, 4.0, thousands(counts["raw rows"]), "rows loaded, every id unique", size=30)
    d.stat(s, 8.7, BODY_Y + 1.45, 4.0, "47 → 13", "columns in the file → columns the questions need", size=30)
    d.stat(s, 8.7, BODY_Y + 2.8, 4.0, "2.7 GB", "loaded in one pass; indexes added after", size=30)
    d.stat(s, 8.7, BODY_Y + 4.15, 4.0, "5 rows", "the routes, as a dbt seed", size=30)

    # 4 — looking at the data
    s = d.content_slide("Looking at the data before answering anything",
                        source=[("query 10", Q[10]), ("query 12", Q[12]), ("data dictionary", L["dict"])])
    d.prose(s, MARGIN, BODY_Y, 7.4, 5.2, [
        "I started with an open-ended sense of curiosity about what was in the data, using the "
        "questions as a framework to look at it through, but mostly just putting my eyes on it: "
        "going through the rows, running quick checks on every field, looking for nulls, common "
        "values and duplicates. That led to several significant findings right off the bat.",
        "There were a ton of duplicate rows. They all looked like the same construction event at "
        "the same time, pulled in again with a different value for something like the weather "
        "reading on a different day. New York City was split by borough instead of combined into "
        "one city. There were numerous batch closures, where thousands of events share the exact "
        "same timestamp; jobs still open when the data was pulled; and seasonal and winter "
        "closures that don't really qualify as construction.",
        "One more thing worth knowing: “hours” in this data is the time an event was "
        "reported as affecting traffic, not hours of actual work. Everything that follows is "
        "hours of reported impact.",
    ], size=13.5)
    d.text(s, 8.7, BODY_Y, 4.0, 0.4, ["What I found, by the numbers"], size=14, bold=True)
    d.table(s, 8.7, BODY_Y + 0.5, [1.9, 2.2], [
        ["Duplicate rows", f"{thousands(dup)}  ({dup / counts['raw rows'] * 100:.1f}%)"],
        ["Seasonal closures", thousands(counts["seasonal closures"])],
        ["Batch start stamps", f"{thousands(con['batch_start'])}  ({n(con['batch_start'])/n(con['events'])*100:.1f}%)"],
        ["Batch end stamps", f"{thousands(con['batch_end'])}  ({n(con['batch_end'])/n(con['events'])*100:.1f}%)"],
        ["Still open at extract", thousands(con["open_at_extract"])],
        ["Typical event", f"median {n(con['median_hours']):.0f} h · mean {con['mean_hours']} h"],
    ], size=11.5, row_h=0.42, header=False)
    d.caption(s, 8.7, BODY_Y + 0.5 + 6 * 0.42 + 0.08, 4.0,
              "Batch stamp: a start or end timestamp shared by 100 or more rows.")

    # 5 — staging
    s = d.content_slide("Staging the data", source=[("staging model", L["staging"]), ("query 13", Q[13])])
    d.prose(s, MARGIN, BODY_Y, 5.6, 5.2, [
        [("Knowing this, I ", {}), ("staged the data", {"link": link(L["staging"])}),
         (": deduped it, combined New York's boroughs into one city, and flagged the seasonal "
          "road closures, the events still open when the data was extracted, and the suspected "
          "administrative batch opens and closes. I flagged rather than deleted, so I could run "
          "any answer with or without them.", {})],
        "In doing so I created a fun error. My first filter on the description field was too "
        "broad: it matched anything with “season” or “winter” in it, which pulled in "
        "things like Seasons Road and Winter Street, and a landmark on CA-108 called the Winter "
        "Closure Gate. I had to refine it so my cleanup wasn't creating noise of its own. The "
        "table shows what it caught.",
        "One thing the numbers told me at this point: the median event is 7 hours and the mean "
        "is about 400. The totals are carried by a long tail of multi-day events, not the typical "
        "one.",
    ], size=13)
    fp = [r for r in examples if r["kind"].startswith("matched")][:5]
    sc = [r for r in examples if r["kind"] == "seasonal closure"][:4]
    rows = [["Caught by the first filter, not a closure", "events"]] + [[r["description"][:88], r["events"]] for r in fp]
    d.table(s, 6.7, BODY_Y, [5.3, 0.7], rows, size=9.5, row_h=0.36, header_size=10)
    y2 = BODY_Y + 0.36 * len(rows) + 0.2
    rows2 = [["Real seasonal closures, kept flagged", "events"]] + [[r["description"][:88], r["events"]] for r in sc]
    d.table(s, 6.7, y2, [5.3, 0.7], rows2, size=9.5, row_h=0.36, header_size=10)

    # 6 — Q1
    s = d.content_slide("Question 1: over the past five years, which city has had the most hours of construction?",
                        source=[("query 1", Q[1]), ("int_city_month_hours", L["int"])])
    d.figure(s, "q1_ten_cities.png", MARGIN, BODY_Y - 0.05, w=7.9)
    d.caption(s, MARGIN, 5.75, 7.9, "All cities in the data, 2017–2021, deduplicated, seasonal closures excluded. Hours are allocated to the months each event was active.")
    d.prose(s, 8.9, BODY_Y, 3.8, 5.2, [
        "Once the data was staged I could actually dig into the questions.",
        f"The top ten by hours is a weird mix. {first['city']} comes out on top, then "
        f"{second['city']}, and both are far ahead of the rest. Los Angeles is notably missing, "
        "and seeing a city like York, Pennsylvania in the top ten seemed a little funny.",
        f"{first['city']} leads only because I combined the boroughs, and if I drop the "
        f"batch-stamped rows it is nearly a tie: {millions(first['hours_excl_batch'])} hours to "
        f"{millions(second['hours_excl_batch'])}. Both of those choices are mine, and both are stated.",
    ], size=12.5)

    # 7 — Q2 by city
    s = d.content_slide("Question 2: how has this changed over time? (1 of 2: the ten cities above)",
                        source=[("query 3", Q[3])])
    pic = d.figure(s, "q2_city_small_multiples.png", MARGIN, BODY_Y - 0.1, h=3.6)
    d.prose(s, MARGIN, (pic.top + pic.height) / 914400 + 0.15, 12.1, 1.9, [
        "I wasn't looking to answer question two yet; it just happened to be the next thing. "
        "Before going further I wanted to know why the list looked the way it did, and looking "
        "year over year it became apparent: the number of events explodes in 2021, everywhere I "
        "looked, in cities under ten different state DOTs. That could be a change in how the data "
        "is measured, a change in the inputs, or a change in coverage. I wondered whether the 2021 "
        "infrastructure bill was behind it, but it was signed in November 2021 and the jump starts "
        "in March, so the timing doesn't fit. Something changed in 2021 that significantly "
        "increased the reported number of events everywhere, and that changes the story quite a bit.",
    ], size=12.5)

    # 8 — Q2 all cities
    s = d.content_slide("Question 2: how has this changed over time? (2 of 2: all cities)",
                        source=[("query 4", Q[4]), ("query 5", Q[5])])
    d.figure(s, "q2_national_monthly.png", MARGIN, BODY_Y - 0.1, w=7.6)
    d.prose(s, 8.5, BODY_Y, 4.3, 2.2, [
        "Looking at all cities together, construction events per month increased significantly, "
        "and construction hours per month increased along with them. The hours fall off at the "
        "very end of 2021, but that isn't real: the jobs still running when the data was pulled "
        "haven't had their end dates recorded yet, so they don't count toward hours in those months.",
    ], size=12)
    d.table(s, 8.5, 3.75, [0.65, 0.95, 0.85, 0.85, 0.9], [
        ["Year", "Events", "Median h", "% = 96 h", "% < 1 h"]] + [
        [r["year"], compact(r["events"]), f"{n(r['median_hours']):.0f}", r["pct_96h"], r["pct_sub_hour"]] for r in decomp
    ], size=10.5, row_h=0.29, header_size=10.5)
    d.prose(s, 8.5, 5.85, 4.3, 1.0, [
        "The spring 2021 jump in hours came from many more long jobs, 30 days and up; the jump "
        "in event counts came from short ones, under an hour. Two different feeds arriving in "
        "the same year.",
    ], size=11)
    d.caption(s, MARGIN, 5.0, 7.6, "All cities, by month. 96-hour events are a provider default (exactly four days); sub-hour events are “slow traffic due to roadwork” observations.")

    # 9 — Q3
    nyc = next(r for r in cmp_rows if r["city"] == "New York")
    sd = next(r for r in cmp_rows if r["city"] == "San Diego")
    s = d.content_slide("Question 3: accounting for varying city sizes, how do the ten largest US cities compare?",
                        source=[("query 7", Q[7]), ("fct_city_comparison", L["cmp_fct"])])
    d.figure(s, "q3_rank_bump.png", MARGIN, BODY_Y - 0.1, h=4.3)
    tx = 6.5
    d.prose(s, tx, BODY_Y, 6.2, 2.6, [
        "Question three asks how the ten largest cities compare once you account for size. "
        "Population didn't feel like the right metric; for a truck the right one is miles of "
        "road, which I couldn't get in the time I had, so I used land area and showed population "
        "alongside it.",
        f"It barely moved the ranking. New York, Phoenix and Chicago are the top three any way you "
        f"cut it, and San Diego is last. Phoenix moves to first per resident. What stands out is "
        f"the spread: New York has about {n(nyc['hours_per_sq_mi'])/n(sd['hours_per_sq_mi']):.0f} times "
        f"San Diego's hours per square mile.",
    ], size=12)
    d.table(s, tx, 4.15, [1.5, 1.1, 1.3, 1.5], [
        ["City", "Hours", "per sq mile", "per 100k people"]] + [
        [r["city"], millions(r["hours"]), compact(r["hours_per_sq_mi"]), compact(r["hours_per_100k"])] for r in cmp_rows[:5]
    ], size=10.5, row_h=0.27, header_size=10.5)
    d.prose(s, tx, 5.85, 6.2, 0.9, [
        "I don't think any two big cities differ eighty-fold in actual roadwork. I think it is how "
        "much each state publishes to these feeds. That is a hypothesis; I can't prove it from "
        "this data alone, and the last slide says what would.",
    ], size=11)
    d.caption(s, MARGIN, 5.75, 5.6, "Land area and population: 2020 Census, city proper. Top five of the ten shown in the table; the chart has all ten.")

    # 10 — Q4
    rank_by_year = {}
    for r in routes_year:
        rank_by_year.setdefault(r["year"], {})[int(r["rank_hours_per_mile"])] = route_name(r)
    top_route = routes_all[0]
    top21 = next(r for r in routes_year if r["path_name"] == top_route["path_name"] and r["year"] == "2021")
    crossing = {}
    for b in bridges:
        if b["crosses"] == "True":
            crossing[route_name(b)] = (b["landmark"].split(" (")[0], b["closest_m"])
    s = d.content_slide("Question 4: which route poses the greatest risk because of construction?",
                        source=[("query 8", Q[8]), ("route matching", L["routes_int"]), ("bridge check", L["bridges"])])
    d.figure(s, "q4_routes.png", MARGIN, BODY_Y - 0.1, w=7.4)
    d.text(s, 9.3, BODY_Y, 3.4, 0.4, ["Leader by year, hours per mile"], size=12.5, bold=True)
    d.table(s, 9.3, BODY_Y + 0.45, [0.7, 1.35, 1.35], [
        ["Year", "1st", "2nd"]] + [
        [y, rank_by_year[y][1], rank_by_year[y][2]] for y in sorted(rank_by_year)
    ], size=11, row_h=0.3, align_right_from=99, header_size=10.5)
    d.prose(s, MARGIN, 4.5, 12.1, 2.2, [
        [("Each route came as a list of about 700 points along the road, a few dozen meters apart, "
          "and each construction event is a point too. So I counted an event toward a route if it "
          "fell within 100 meters of any of the route's points, checking both ends of the event and, "
          "for the short ones, a few points along the line between them. Then I added up the hours "
          "per route and divided by the route's length, so a longer route isn't riskier just for "
          "being long. It is all ", {}), ("SQL in the repo", {"link": link(L["routes_int"])}), (".", {})],
        [(f"{route_name(top_route)} carries the most construction per mile, over the five years and "
          f"in the most recent year, and weighting by severity gives the same order. But the leader "
          f"changed every couple of years, and {route_name(top_route)}'s 2021 is "
          f"{n(top21['hours'])/n(top_route['hours'])*100:.0f}% of its five-year total, the same 2021 "
          f"effect as the cities. To make sense of that I ", {}),
         ("checked where the routes go", {"link": link(L["bridges"])}),
         (f": Berkeley and San Leandro both cross the Bay Bridge (their points pass "
          f"{crossing.get('Berkeley', ('', '?'))[1]} m from Yerba Buena Island), San Rafael crosses "
          f"the Golden Gate ({crossing.get('San Rafael', ('', '?'))[1]} m from mid-span), and San "
          f"Mateo and Pacifica stay on the peninsula. The three bridge routes are the top three every "
          f"year; Pacifica is at or near the bottom every year.", {})],
    ], size=11.5)

    # 11 — Q4 check
    s = d.content_slide("Question 4: checking what the matching actually found",
                        source=[("query 9", Q[9]), ("route matching", L["routes_int"])])
    d.table(s, MARGIN, BODY_Y, [6.6, 0.9, 1.0], [
        ["Description, as published", "Events", "Hours"]] + [
        [r["description"][:95], r["events"], compact(r["hours"])] for r in desc[:11]
    ], size=10.5, row_h=0.36, header_size=10.5)
    d.prose(s, 9.5, BODY_Y, 3.2, 5.0, [
        "To make sure the matching did what I thought, I pulled the biggest events on the "
        "Berkeley route in 2021: I-80 at Ashby, Cleveland Street, the Bay Bridge piers, Treasure "
        "Island, the Yerba Buena tunnel. That is what I'd expect to see on that road.",
        f"The weak spot is long events. An event is two points and the road between them is "
        f"unknown, so anything over five miles is matched on its endpoints only: "
        f"{thousands(counts['bay area events over 5 miles'])} of the "
        f"{thousands(counts['bay area events'])} Bay Area events. San Leandro has the most of "
        f"those and is the route most likely undercounted.",
    ], size=12)
    d.caption(s, MARGIN, BODY_Y + 12 * 0.36 + 0.1, 8.5, "Berkeley route, 2021, by hours matched. The Van Ness and Mission entries are the shared downtown start, which counts for all five routes.")

    # 12 — Q5
    s = d.content_slide("Question 5: how would I present the Q4 answer to engineers versus executives?")
    col_w = 5.85
    for i, (head_, paras) in enumerate([
        ("To an engineering group", [
            "I found this one hard, because there isn't a definitive answer to give: the riskiest "
            "route keeps changing year over year. For engineers that is the answer. I'd show the "
            "year-by-year table, walk through how the matching works, and be specific about where "
            "it is weak: events are two points, long ones get undercounted, two of the routes "
            "share the Bay Bridge approach.",
            "Then what I'd want to make it better: road geometry instead of straight lines, a "
            "provider ID on each event, re-published closures chained into one, and exposure per "
            "road segment rather than per route. And I'd hand over the models, because they will "
            "want to re-run it with their own buffer.",
        ]),
        ("To business and operations executives", [
            "I'd lead with what I can say clearly: Pacifica is the lowest-exposure route every "
            "year, the three routes that cross a bridge carry roughly twice the construction per "
            "mile, and Berkeley leads in the most recent year. If one route runs for a year, "
            "Pacifica minimizes construction exposure and Berkeley maximizes it.",
            "Then one sentence on confidence: high on the bottom of the ranking, moderate on the "
            "top, because the order among the bridge routes shifts. One chart, one number per "
            "route, no table of years. The assumptions on a backup slide, offered if anyone asks. "
            "The uncertainty doesn't go away for executives; it gets compressed into a confidence "
            "statement instead of a table.",
        ]),
    ]):
        x = MARGIN + i * (col_w + 0.4)
        d.rect(s, x, BODY_Y, col_w, 0.55, fill=LIGHT)
        d.text(s, x + 0.2, BODY_Y, col_w - 0.4, 0.55, [head_], size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        d.prose(s, x + 0.2, BODY_Y + 0.75, col_w - 0.4, 4.4, paras, size=12.5)

    # 13 — Q6
    s = d.content_slide("Question 6: what AI tools did I use, how, and why?", source=[("process log", L["log"])])
    d.prose(s, MARGIN, BODY_Y, 5.9, 5.2, [
        "I used ChatGPT to walk me through setting up the environment step by step: a second "
        "Postgres instance, the virtual environment, the repo, dbt. It told me what to run and I "
        "ran it and read the output.",
        "I used Claude for the rest, in two ways. As a consultant, to talk through how to approach "
        "the data and what the problems I was finding meant. And as a code writer, once I had "
        "decided what I wanted: the dbt models, the loader, the queries behind this deck.",
        [("I made the calls. The dedup rule, the words for the seasonal-closure filter, the "
          "five-year window, the 100-meter buffer, which year to lead with. I ran every query, read "
          "every result, and kept a ", {}), ("log of the whole process", {"link": link(L["log"])}),
         (" in the repo.", {})],
    ], size=12.5)
    d.text(s, 7.0, BODY_Y, 5.7, 0.4, ["Where it was wrong"], size=14, bold=True)
    d.prose(s, 7.0, BODY_Y + 0.5, 5.7, 4.7, [
        "It told me my Mac didn't have Postgres installed. It did; the assistant's shell runs in "
        "a Linux VM beside the Mac and couldn't see it. Caught on the first exchange.",
        "It wrote a hypothesis as a finding: that the spread between cities “is” reporting "
        "coverage. The table shows a spread, not its cause. I had it rewritten as a hypothesis with "
        "the tests that would settle it.",
        "The first winter-closure filter was too broad; I gave it the word list and neither of us "
        "anticipated Winter Street. Found by exporting the matches and reading them.",
        "A typical ask, so you can see the shape of it: “We need to dedupe rows that are the "
        "same place with the same timestamps. I want to use the description to separate "
        "construction from winter closures. Thoughts?”",
    ], size=12.5)

    # 14 — assumptions and more time
    s = d.content_slide("What I assumed, and what I'd do with more time")
    d.text(s, MARGIN, BODY_Y, 6.0, 0.4, ["Assumptions I made"], size=14, bold=True)
    d.prose(s, MARGIN, BODY_Y + 0.5, 6.0, 4.8, [
        "A city is the municipality, and New York is its five boroughs.",
        "Rows with the same start, end, location and description are one event.",
        "“The past five years” is the five complete years, 2017 through 2021.",
        "Seasonal closures are not construction, and the description identifies them.",
        "Hours of reported traffic impact stand in for hours of roadwork.",
        "An event is its two endpoints; the road between them is unknown.",
        "“Near a route” is within 100 meters; short events follow a straight line.",
        "Raw counts reflect how much each state publishes as much as how much work there is.",
    ], size=12.5)
    d.text(s, 7.1, BODY_Y, 5.6, 0.4, ["With more time"], size=14, bold=True)
    d.prose(s, 7.1, BODY_Y + 0.5, 5.6, 4.8, [
        "Use miles of road as the denominator for question three, pulled from OpenStreetMap "
        "inside each city's boundary.",
        "Chain re-published closures into one event, so a five-month closure is five months and "
        "not sixty five-hour rows.",
        "Work out which provider each row came from. The wind column is spelled two ways, and that "
        "probably dates the feed changes. Their share by state would test the coverage hypothesis, "
        "along with an outside benchmark like FHWA state DOT spending per lane-mile.",
        "Snap events to road geometry and score exposure per road segment, which is what a "
        "planner could act on.",
        "Or go simpler, not fancier: for a one-off, one notebook reading the CSV directly would "
        "have been faster than a warehouse.",
    ], size=12.5)

    # 15 — appendix
    s = d.content_slide("Appendix: where every number comes from",
                        takeaway="pull_deck_data.py runs these against the database and writes one CSV each; build_charts.py and build_deck.py read the CSVs and nothing else.",
                        source=[("reports/", "reports/README.md"), ("deck builder", L["deck"])])
    d.table(s, MARGIN, BODY_Y, [1.1, 3.6, 4.0, 3.4], [
        ["Slides", "Query", "Reads", "Shows"],
        ["3–5", "queries/10, 12, 13", "raw, stg_construction", "row counts, flag counts, what the first filter caught"],
        ["6", "queries/01_q1_top_cities", "int_city_month_hours", "Q1 top ten by hours"],
        ["7", "queries/03_q2_ten_cities_by_year", "int_city_month_hours", "the same ten by year"],
        ["8", "queries/04, queries/05", "int_city_month_hours, stg_construction", "monthly totals; by-year decomposition"],
        ["9", "queries/07_q3_city_comparison", "fct_city_comparison", "Q3 normalized ranking"],
        ["10", "queries/08, check_route_bridges.py", "fct_route_exposure, sf_paths", "Q4 by route and year; which routes cross a bridge"],
        ["11", "queries/09_q4_berkeley_2021_descriptions", "int_route_events", "matched descriptions"],
        ["—", "queries/02, 06, 11", "various", "log only; not on a slide"],
    ], size=11, row_h=0.38, align_right_from=99, header_size=11)

    d.save()


if __name__ == "__main__":
    build()
